import pandas as pd
from scipy import stats
from docx import Document
from docx.shared import Inches

# Load the data from the CSV file
data = pd.read_csv("GNG4120.csv")

# Drop the 'Student ID' column
data.drop(columns=['Student ID'], inplace=True)

# Calculate mean and standard deviation for each question
mean_values = data.mean()
std_dev_values = data.std()

# Perform a chi-square test for each pair of questions (assuming they are categorical variables)
chi_square_results = pd.DataFrame(index=data.columns, columns=data.columns)

for col1 in data.columns:
    for col2 in data.columns:
        contingency_table = pd.crosstab(data[col1], data[col2])
        chi2, p, _, _ = stats.chi2_contingency(contingency_table)
        chi_square_results.loc[col1, col2] = p

# Create a Word document
doc = Document()
doc.add_heading('Statistical Analysis', level=1)

# Add mean and standard deviation table
doc.add_heading('Mean and Standard Deviation', level=2)
mean_std_table = doc.add_table(rows=len(mean_values)+1, cols=3)
hdr_cells = mean_std_table.rows[0].cells
hdr_cells[0].text = 'Question'
hdr_cells[1].text = 'Mean'
hdr_cells[2].text = 'Standard Deviation'

for i, (question, mean_val, std_dev_val) in enumerate(zip(mean_values.index, mean_values, std_dev_values)):
    row_cells = mean_std_table.rows[i+1].cells
    row_cells[0].text = question
    row_cells[1].text = f"{mean_val:.2f}"
    row_cells[2].text = f"{std_dev_val:.2f}"

# Add chi-square p-values table
doc.add_heading('Chi-square p-values', level=2)
chi_square_table = doc.add_table(rows=len(data.columns)+1, cols=len(data.columns)+1)
hdr_cells = chi_square_table.rows[0].cells
for i, col in enumerate(data.columns):
    hdr_cells[i+1].text = col
for i, (index, row) in enumerate(chi_square_results.iterrows()):
    row_cells = chi_square_table.rows[i+1].cells
    row_cells[0].text = index
    for j, p_value in enumerate(row):
        row_cells[j+1].text = f"{p_value:.4f}"

# Save the document
doc.save('statistical_analysis.docx')
